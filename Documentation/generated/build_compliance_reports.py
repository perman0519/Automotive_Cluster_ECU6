from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
REQ_DOC = ROOT / "MPC5748G_Embedded_Training_Requirements_EN.docx"
STATUS_MD = ROOT / "Documentation" / "ecu6_gateway" / "requirements_status.md"
EVIDENCE_MD = ROOT / "Documentation" / "ecu6_gateway" / "evidence" / "evidence_index.md"
TEST_REPORT_MD = ROOT / "Documentation" / "ecu6_gateway" / "test_report.md"
HOST_TEST_LOG = ROOT / "Documentation" / "ecu6_gateway" / "evidence" / "host_unit_tests_2026-06-04.txt"
OUT_DIR = ROOT / "Documentation" / "generated"
FONT_KR = Path("C:/Windows/Fonts/malgun.ttf")
FONT_KR_BOLD = Path("C:/Windows/Fonts/malgunbd.ttf")
FONT_EN = Path("C:/Windows/Fonts/calibri.ttf")
FONT_EN_BOLD = Path("C:/Windows/Fonts/calibrib.ttf")


COMMON_EVIDENCE = {
    "REQ-COM-001": {
        "impl": "Sources/main.c; Sources/services/system_init.c; Sources/drivers/gpio.c; Sources/drivers/can.c",
        "verify": "Source inspection; Debug_FLASH artifact inspection; pending S32DS runtime/flash evidence",
        "evidence": "Debug_FLASH/Automotive_Cluster.elf; Debug_FLASH/Automotive_Cluster.map; cmm/interrupt_evidence.cmm",
        "note": "Runtime stability, UART boot log, heartbeat LED, and Trace32 screenshot evidence are still required.",
    },
    "REQ-COM-002": {
        "impl": "Sources/services/logger.c; Sources/automotive.h; Sources/services/com_gateway.c",
        "verify": "Source inspection of LOG_Print() format and gateway/OTA logging call sites",
        "evidence": "Sources/services/logger.c; Sources/services/com_gateway.c",
        "note": "Saved UART terminal evidence was not found, but code implements the required format.",
    },
    "REQ-COM-003": {
        "impl": "Sources/drivers/can.c; Sources/services/com_gateway.c; CANOE/*.dbc; CANOE/*.can",
        "verify": "Source/DBC/CAPL inspection; readable CANoe verdict export still pending",
        "evidence": "CANOE/6_ECUs_dbc_0206.dbc; CANOE/ECU6_Gateway_Test.can; CANOE/*.vtestreport",
        "note": "Binary CANoe reports exist, but text-readable PASS/FAIL verdicts were not available.",
    },
    "REQ-COM-004": {
        "impl": "Sources/services/ms_scheduler.c; Sources/main.c; cmm/scheduler_evidence.cmm",
        "verify": "Source inspection; Trace32 scheduler counter capture pending",
        "evidence": "Sources/services/ms_scheduler.c; cmm/scheduler_evidence.cmm; Debug_FLASH/Automotive_Cluster.map",
        "note": "Scheduler counters are present in source; a rebuild and Trace32 evidence are required for final runtime closure.",
    },
    "REQ-COM-005": {
        "impl": "Sources/services/com_gateway.c: COM_GatewayCheckTimeouts(), COM_GatewayHandleRxMessage()",
        "verify": "Source inspection and host unit test coverage",
        "evidence": "Tests/host/test_ecu6_gateway_logic.py; Documentation/ecu6_gateway/evidence/host_unit_tests_2026-06-04.txt",
        "note": "Host tests passed; hardware/CANoe runtime evidence remains recommended.",
    },
    "REQ-COM-006": {
        "impl": "Sources/services/com_gateway.c: COM_GatewayBuildStatusMessage(); CANOE/6_ECUs_dbc_0206.dbc",
        "verify": "Source and DBC inspection",
        "evidence": "Sources/services/com_gateway.c; CANOE/6_ECUs_dbc_0206.dbc",
        "note": "GatewayStatus byte 6 carries the 4-bit alive counter.",
    },
    "REQ-COM-007": {
        "impl": "Sources/services/com_gateway.c: COM_CalculateChecksum(), COM_VerifyChecksum()",
        "verify": "Source inspection and host unit test coverage",
        "evidence": "Tests/host/test_ecu6_gateway_logic.py; host_unit_tests_2026-06-04.txt",
        "note": "Host tests passed for checksum calculation and validation.",
    },
    "REQ-COM-008": {
        "impl": "Sources/main.c; Sources/drivers/*.c; Sources/services/*.c; Generated_Code/*.c",
        "verify": "Source tree review",
        "evidence": "Sources/drivers; Sources/services; Generated_Code",
        "note": "Driver/service/generated layers are separated.",
    },
    "REQ-COM-009": {
        "impl": "Sources/services/com_gateway.c: gateway, OTA, and security state handling",
        "verify": "Source inspection and host unit test coverage",
        "evidence": "Documentation/ecu6_gateway/design.md; Tests/host/test_ecu6_gateway_logic.py",
        "note": "Gateway, OTA, and security modes are explicitly represented.",
    },
    "REQ-COM-010": {
        "impl": "cmm/init_debug.cmm; cmm/test_debug.cmm; cmm/*.cmm",
        "verify": "File inspection",
        "evidence": "cmm/init_debug.cmm; cmm/test_debug.cmm; cmm/interrupt_evidence.cmm",
        "note": "Trace32 screenshots are still required for runtime evidence.",
    },
    "REQ-COM-011": {
        "impl": "CANOE/ECU6_Gateway_Test.can; CANOE/ECU6_OTA_State_Test.can",
        "verify": "CAPL file inspection; binary report artifacts observed",
        "evidence": "CANOE/ECU6_Gateway_Test.can; CANOE/ECU6_OTA_State_Test.can; CANOE/*.vtestreport",
        "note": "Readable CANoe PASS/FAIL export is pending.",
    },
    "REQ-COM-012": {
        "impl": "Documentation/ecu6_gateway/design.md; Documentation/ecu6_gateway/test_report.md",
        "verify": "Document inspection",
        "evidence": "Documentation/ecu6_gateway/design.md; Documentation/ecu6_gateway/test_report.md",
        "note": "Design and test documentation exists.",
    },
    "REQ-COM-013": {
        "impl": "Sources/services/com_gateway.c: timeout/alive/checksum recovery paths",
        "verify": "Source inspection and host unit test coverage",
        "evidence": "Tests/host/test_ecu6_gateway_logic.py; Documentation/ecu6_gateway/design.md",
        "note": "Host tests include alive and timeout recovery.",
    },
    "REQ-COM-014": {
        "impl": "Sources/services/com_gateway.c comments; cmm/gateway_debug.cmm; Documentation/ecu6_gateway/*.md",
        "verify": "Traceability review",
        "evidence": "Sources/services/com_gateway.c; cmm/gateway_debug.cmm; Documentation/ecu6_gateway/requirements_status.md",
        "note": "Requirement IDs appear in code, scripts, and status documents.",
    },
    "REQ-COM-015": {
        "impl": "Tests/host/test_ecu6_gateway_logic.py",
        "verify": "Executed Python unittest discovery",
        "evidence": "Documentation/ecu6_gateway/evidence/host_unit_tests_2026-06-04.txt",
        "note": "9/9 host-side unit tests passed.",
    },
}

ECU6_EVIDENCE = {
    "REQ-ECU6-001": ("Sources/drivers/can.c; Sources/services/com_gateway.c", "Source/DBC/CAPL inspection", "CANOE/6_ECUs_dbc_0206.dbc; CANOE/ECU6_Gateway_Test.can", "Allowed IDs 0x101, 0x201, 0x301, 0x401, and 0x501 are monitored."),
    "REQ-ECU6-002": ("Sources/drivers/can.c; Sources/services/com_gateway.c: COM_GatewayHandleRxMessage()", "Source inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Unknown ID increments blocked counter."),
    "REQ-ECU6-003": ("Sources/services/com_gateway.c: COM_VerifyChecksum()", "Source inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Checksum mismatch sets error/counting behavior."),
    "REQ-ECU6-004": ("Sources/services/com_gateway.c: COM_UpdateAliveMonitor()", "Source inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Alive freeze detection and recovery tested on host."),
    "REQ-ECU6-005": ("Sources/services/com_gateway.c: OTA_ApplyCommandState(); CANOE/6_ECUs_dbc_0206.dbc", "Source/DBC/CAPL inspection", "CANOE/ECU6_OTA_State_Test.can", "OTA states are implemented and exposed in GatewayStatus."),
    "REQ-ECU6-006": ("Sources/drivers/can.c; Sources/services/com_gateway.c: OTA_HandleRequestMessage()", "Source/DBC/CAPL inspection", "CANOE/6_ECUs_dbc_0206.dbc; CANOE/ECU6_OTA_State_Test.can", "0x650 request and 0x651 response are implemented."),
    "REQ-ECU6-007": ("Sources/services/com_gateway.c: OTA_ApplyCommandState()", "Source/CAPL inspection", "CANOE/ECU6_OTA_State_Test.can", "START/DATA/END/ACTIVATE/ABORT paths and invalid sequence NACK are implemented."),
    "REQ-ECU6-008": ("Sources/services/com_gateway.c: OTA_StoreDataPayload()", "Source/CAPL inspection", "CANOE/ECU6_OTA_State_Test.can", "Bounded 256-byte buffer and overflow rejection are implemented."),
    "REQ-ECU6-009": ("Sources/services/com_gateway.c: OTA_VerifyDataChecksum()", "Source inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Payload XOR checksum success/failure is tested."),
    "REQ-ECU6-010": ("Sources/services/com_gateway.c: OTA_ApplyCommandState()", "Source inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "START is rejected before security unlock."),
    "REQ-ECU6-011": ("Sources/services/com_gateway.c: Security_GenerateSeed(), Security_CalculateKey()", "Source inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Seed/key unlock and wrong-key handling are tested."),
    "REQ-ECU6-012": ("Sources/services/com_gateway.c: OTA_VerifyDataAuth()", "Source/CAPL inspection", "CANOE/ECU6_OTA_State_Test.can; cmm/ota_debug.cmm", "OTA_DATA AuthByte rejection is implemented."),
    "REQ-ECU6-013": ("Sources/services/com_gateway.c: Security_RecordFailure(), Security_IsLockoutActive()", "Source inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Three failed attempts trigger lockout and recovery timeout."),
    "REQ-ECU6-014": ("Sources/services/com_gateway.c: COM_GatewayBuildStatusMessage(); Sources/drivers/can.c: vCANTask()", "Source/DBC inspection", "Debug_FLASH/Automotive_Cluster.map; CANOE/6_ECUs_dbc_0206.dbc", "GatewayStatus 0x601 is built and scheduled every 100 ms."),
    "REQ-ECU6-015": ("Sources/services/com_gateway.c: OTA_AdvanceActiveSwVersion()", "Source inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Version increments only after verified activation."),
    "REQ-ECU6-016": ("cmm/gateway_debug.cmm; cmm/ota_debug.cmm; cmm/security_debug.cmm", "File inspection", "cmm/*.cmm", "Trace32 scripts watch gateway, OTA, and security variables."),
}

INT_EVIDENCE = {
    "REQ-INT-001": ("CANOE/ECU1.can..ECU5.can; Sources/drivers/can.c; Sources/services/com_gateway.c; Configuration1.cfg", "Static integration setup inspection; full CANoe run pending", "Configuration1.cfg; CANOE/6_ECUs_dbc_0206.dbc; CANOE/*.can", "Full six-board CANoe/bench trace is still required."),
    "REQ-INT-007": ("Sources/services/com_gateway.c; CANOE/ECU6_OTA_State_Test.can", "Source/CAPL inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "OTA success logic, activation, and version update are covered by host tests."),
    "REQ-INT-008": ("Sources/services/com_gateway.c; CANOE/ECU6_OTA_State_Test.can", "Source/CAPL inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Wrong checksum failure keeps version unchanged."),
    "REQ-INT-009": ("Sources/services/com_gateway.c; CANOE/ECU6_Gateway_Test.can", "Source/CAPL inspection and host unit test", "Tests/host/test_ecu6_gateway_logic.py", "Invalid ID and checksum error behavior is implemented and host-tested."),
}

ACTIVE_EVIDENCE = {**COMMON_EVIDENCE}
for req_id, row in ECU6_EVIDENCE.items():
    ACTIVE_EVIDENCE[req_id] = {"impl": row[0], "verify": row[1], "evidence": row[2], "note": row[3]}
for req_id, row in INT_EVIDENCE.items():
    ACTIVE_EVIDENCE[req_id] = {"impl": row[0], "verify": row[1], "evidence": row[2], "note": row[3]}


KO_REQ_DESC = {
    "REQ-COM-001": "각 ECU는 시스템 클록, GPIO, UART, CAN 컨트롤러, 타이머/스케줄러 및 필요한 인터럽트를 초기화해야 한다.",
    "REQ-COM-002": "각 ECU는 [time][ECUx][module][level] message 형식의 일관된 UART 로그를 제공해야 한다.",
    "REQ-COM-003": "각 ECU는 최소 하나의 주기 CAN 메시지를 송신하고 다른 ECU 또는 CANoe의 CAN 메시지를 수신해야 한다.",
    "REQ-COM-004": "각 ECU는 10 ms, 100 ms, 1000 ms 작업 또는 동등한 스케줄러를 구현해야 한다.",
    "REQ-COM-005": "각 ECU는 중요 수신 CAN 메시지 또는 외부 센서 입력의 timeout 감시를 구현해야 한다.",
    "REQ-COM-006": "각 ECU의 primary status 메시지는 4-bit alive counter를 포함해야 한다.",
    "REQ-COM-007": "각 ECU의 primary status 메시지는 단순 checksum을 포함해야 한다.",
    "REQ-COM-008": "각 ECU는 driver, service, application 계층을 분리해야 한다.",
    "REQ-COM-009": "여러 동작 모드가 있는 ECU는 상태 기계로 동작을 구현해야 한다.",
    "REQ-COM-010": "각 멘티는 Trace32용 init_debug.cmm 및 test_debug.cmm을 작성해야 한다.",
    "REQ-COM-011": "각 멘티는 담당 ECU에 대해 CAPL 테스트 케이스를 최소 3개 작성해야 한다.",
    "REQ-COM-012": "각 멘티는 설계 문서와 테스트 문서를 제공해야 한다.",
    "REQ-COM-013": "각 ECU는 입력 또는 메시지가 정상 상태로 복귀할 때 fault recovery를 지원해야 한다.",
    "REQ-COM-014": "각 ECU는 중요 로직을 요구사항 ID와 연결할 수 있는 traceable name/comment를 사용해야 한다.",
    "REQ-COM-015": "각 ECU는 checksum, filter, warning rule, state transition 등 순수 로직에 대한 host-side unit test를 포함할 수 있다.",
}


def ko_description(req_id: str, statement: str) -> str:
    if req_id in KO_REQ_DESC:
        return KO_REQ_DESC[req_id]
    prefix = req_id.split("-")[1]
    s = statement
    replacements = [
        ("MUST", "반드시"),
        ("SHOULD", "권장 수준으로"),
        ("MAY", "선택적으로"),
        ("receive or read", "수신 또는 읽기"),
        ("read or receive", "읽기 또는 수신"),
        ("send", "송신"),
        ("support", "지원"),
        ("provide", "제공"),
        ("detect", "검출"),
        ("validate", "검증"),
        ("implement", "구현"),
        ("every", "주기로"),
        ("scenario", "시나리오"),
        ("CAN ID", "CAN ID"),
    ]
    for old, new in replacements:
        s = s.replace(old, new)
    if prefix.startswith("ECU"):
        return f"{prefix} 범위 요구사항 해석: {s}"
    if prefix == "INT":
        return f"통합 요구사항 해석: {s}"
    return f"요구사항 해석: {s}"


def load_requirements():
    doc = Document(REQ_DOC)
    rows = []
    for ti, table in enumerate(doc.tables):
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if cells and re.match(r"^REQ-[A-Z0-9]+-\d+", cells[0]):
                rows.append(
                    {
                        "id": cells[0],
                        "priority": cells[1] if len(cells) > 1 else "",
                        "statement": cells[2] if len(cells) > 2 else "",
                        "acceptance": cells[3] if len(cells) > 3 else "",
                        "evidence_required": cells[4] if len(cells) > 4 else "",
                        "table": ti,
                    }
                )
    return rows


def load_statuses():
    text = STATUS_MD.read_text(encoding="utf-8")
    statuses = {}
    for line in text.splitlines():
        m = re.match(r"\|\s*(REQ-[A-Z0-9]+-\d+)\s*\|\s*(Complete|Partial|Missing)\s*\|\s*(.*?)\s*\|", line)
        if m:
            statuses[m.group(1)] = {"raw": m.group(2), "summary": m.group(3)}
    return statuses


def status_for(req_id, statuses, lang):
    raw = statuses.get(req_id, {}).get("raw")
    if raw == "Complete":
        return "충족" if lang == "ko" else "Compliant"
    if raw == "Partial":
        return "부분 충족" if lang == "ko" else "Partially Compliant"
    if raw == "Missing":
        return "미충족" if lang == "ko" else "Non-compliant"
    return "확인 필요" if lang == "ko" else "Needs Confirmation"


def section_for(req_id):
    if req_id.startswith("REQ-COM"):
        return "Common"
    if req_id.startswith("REQ-ECU1"):
        return "ECU1 Body Control"
    if req_id.startswith("REQ-ECU2"):
        return "ECU2 Powertrain"
    if req_id.startswith("REQ-ECU3"):
        return "ECU3 Sensor Fusion"
    if req_id.startswith("REQ-ECU4"):
        return "ECU4 Cluster Display"
    if req_id.startswith("REQ-ECU5"):
        return "ECU5 Diagnostic/UDS"
    if req_id.startswith("REQ-ECU6"):
        return "ECU6 Gateway/OTA/Security"
    return "Integration"


def get_evidence(req_id):
    if req_id in ACTIVE_EVIDENCE:
        return ACTIVE_EVIDENCE[req_id]
    prefix = req_id.split("-")[1]
    if prefix in {"ECU1", "ECU2", "ECU3", "ECU4", "ECU5"}:
        capl = ROOT / "CANOE" / f"{prefix}.can"
        return {
            "impl": f"No {prefix} MPC5748G board source implementation found in this workspace; supporting CAPL simulator: {capl.relative_to(ROOT) if capl.exists() else 'not found'}",
            "verify": "Workspace inspection only; board implementation and execution evidence required",
            "evidence": "MPC5748G_Embedded_Training_Requirements_EN.docx; CANOE support files where present",
            "note": "Insufficient evidence for the assigned ECU board implementation.",
        }
    return {
        "impl": "No complete full-system implementation evidence found in this workspace",
        "verify": "Workspace inspection only; full CANoe/bench execution evidence required",
        "evidence": "Configuration1.cfg; CANOE/*.dbc; CANOE/*.can",
        "note": "Full integration execution or exported CANoe report is required.",
    }


def set_cell_text(cell, text, font="Calibri", size=8.0, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "bottom": bottom, "start": start, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(widths)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        tbl_grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths, lang="en", font_size=8.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for idx, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], h, size=font_size, bold=True)
        shade_cell(table.rows[0].cells[idx], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for idx, val in enumerate(row):
            set_cell_text(cells[idx], val, size=font_size)
    set_table_widths(table, widths)
    doc.add_paragraph()
    return table


def set_doc_styles(doc, lang):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.10


def add_para(doc, text, bold=False, size=None, color=None, align=None, style=None, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    r.bold = bold
    return p


def add_cover(doc, lang):
    title = (
        "Automotive_Cluster 프로젝트 요구사항 충족 증명 보고서"
        if lang == "ko"
        else "Automotive_Cluster Requirements Compliance Evidence Report"
    )
    subtitle = (
        "요구사항-구현-검증-증거 추적성 기반 공식 기술 검증 보고서"
        if lang == "ko"
        else "Formal technical verification report based on requirement-to-evidence traceability"
    )
    date_label = datetime.now().strftime("%Y-%m-%d")
    add_para(doc, "Compliance Evidence Report" if lang == "en" else "요구사항 충족 증명 보고서", bold=True, size=11, color="555555", after=10)
    add_para(doc, title, bold=True, size=24, color="0B2545", after=6)
    add_para(doc, subtitle, size=13, color="555555", after=18)
    rows = [
        ("Project" if lang == "en" else "프로젝트", "Automotive_Cluster"),
        ("Workspace" if lang == "en" else "워크스페이스", str(ROOT)),
        ("Source requirement document" if lang == "en" else "원 요구사항 문서", "MPC5748G_Embedded_Training_Requirements_EN.docx"),
        ("Prepared date" if lang == "en" else "작성일", date_label),
        ("Verification basis" if lang == "en" else "검증 기준", "Inspection, static artifacts, host unit tests, and available logs"),
    ]
    add_table(doc, ["Field", "Value"] if lang == "en" else ["항목", "값"], rows, [2300, 7060], lang, 9.5)
    doc.add_page_break()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        r.font.name = "Calibri"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")


def font(path: Path, size: int, fallback: Path | None = None):
    chosen = path if path.exists() else (fallback if fallback and fallback.exists() else FONT_EN)
    return ImageFont.truetype(str(chosen), size)


def wrap_lines(draw, text, font_obj, max_width):
    lines = []
    for raw_line in str(text).split("\n"):
        words = raw_line.split()
        if not words:
            lines.append("")
            continue
        line = ""
        for word in words:
            candidate = word if not line else f"{line} {word}"
            width = draw.textbbox((0, 0), candidate, font=font_obj)[2]
            if width <= max_width:
                line = candidate
            else:
                if line:
                    lines.append(line)
                line = word
        if line:
            lines.append(line)
    return lines


def centered_text(draw, box, text, font_obj, fill="#0f172a", spacing=8):
    x0, y0, x1, y1 = box
    lines = wrap_lines(draw, text, font_obj, x1 - x0 - 28)
    heights = []
    for line in lines:
        bb = draw.textbbox((0, 0), line, font=font_obj)
        heights.append(bb[3] - bb[1])
    total_h = sum(heights) + spacing * max(0, len(lines) - 1)
    y = y0 + ((y1 - y0) - total_h) / 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=font_obj)[2]
        draw.text((x0 + ((x1 - x0) - width) / 2, y), line, font=font_obj, fill=fill)
        y += height + spacing


def label(draw, xy, text, font_obj, fill="#334155"):
    draw.text(xy, text, font=font_obj, fill=fill)


def rounded(draw, box, text, font_obj, fill="#ffffff", outline="#94a3b8", text_fill="#0f172a", width=3):
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=width)
    centered_text(draw, box, text, font_obj, text_fill)


def arrow(draw, start, end, fill="#2563eb", width=5):
    draw.line([start, end], fill=fill, width=width)
    x0, y0 = start
    x1, y1 = end
    if abs(x1 - x0) >= abs(y1 - y0):
        if x1 >= x0:
            pts = [(x1, y1), (x1 - 18, y1 - 10), (x1 - 18, y1 + 10)]
        else:
            pts = [(x1, y1), (x1 + 18, y1 - 10), (x1 + 18, y1 + 10)]
    else:
        if y1 >= y0:
            pts = [(x1, y1), (x1 - 10, y1 - 18), (x1 + 10, y1 - 18)]
        else:
            pts = [(x1, y1), (x1 - 10, y1 + 18), (x1 + 10, y1 + 18)]
    draw.polygon(pts, fill=fill)


def make_system_architecture(lang: str):
    is_ko = lang == "ko"
    path = OUT_DIR / f"compliance_system_architecture_{lang}.png"
    img = Image.new("RGB", (2200, 1350), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = font(FONT_KR_BOLD if is_ko else FONT_EN_BOLD, 56, FONT_EN_BOLD)
    body_font = font(FONT_KR if is_ko else FONT_EN, 31, FONT_EN)
    small_font = font(FONT_KR if is_ko else FONT_EN, 25, FONT_EN)
    title = "Automotive_Cluster 시스템 아키텍처" if is_ko else "Automotive_Cluster System Architecture"
    label(draw, (70, 50), title, title_font, "#0b2545")

    bus = (410, 565, 1235, 660)
    rounded(draw, bus, "Shared CAN Bus / VN5650", body_font, fill="#dbeafe", outline="#2563eb")
    ecu_labels = [
        ("ECU1\nBodyStatus\n0x101", (110, 280, 390, 420)),
        ("ECU2\nPowertrain\n0x201", (110, 720, 390, 860)),
        ("ECU3\nSensorStatus\n0x301", (540, 280, 820, 420)),
        ("ECU4\nClusterStatus\n0x401", (540, 720, 820, 860)),
        ("ECU5\nDiagnostic\n0x501", (970, 280, 1250, 420)),
        ("ECU6\nGateway/OTA/Security\n0x601, 0x650, 0x651", (1390, 280, 1770, 860)),
    ]
    for text, box in ecu_labels:
        fill = "#dcfce7" if "ECU6" not in text else "#fef3c7"
        outline = "#22c55e" if "ECU6" not in text else "#f59e0b"
        rounded(draw, box, text, body_font if "ECU6" not in text else small_font, fill=fill, outline=outline)
        if "ECU6" in text:
            continue
        arrow(draw, ((box[0] + box[2]) // 2, box[3]), ((box[0] + box[2]) // 2, bus[1]), "#64748b", 4)
        if box[1] > bus[3]:
            arrow(draw, ((box[0] + box[2]) // 2, box[1]), ((box[0] + box[2]) // 2, bus[3]), "#64748b", 4)
    arrow(draw, (1235, 612), (1390, 612), "#2563eb", 5)

    rounded(draw, (1850, 280, 2150, 430), "CANoe\nDBC/CAPL/Panel\nReports", small_font, fill="#eef2ff", outline="#6366f1")
    arrow(draw, (1850, 590), (1810, 590), "#6366f1", 5)
    rounded(draw, (1850, 520, 2150, 670), "Trace32\nCMM scripts\nBreakpoints", small_font, fill="#f1f5f9", outline="#475569")
    arrow(draw, (1850, 595), (1770, 595), "#475569", 5)
    rounded(draw, (1850, 760, 2150, 910), "S32DS\nDebug_FLASH\nELF/MAP", small_font, fill="#fee2e2", outline="#ef4444")
    arrow(draw, (1850, 835), (1770, 835), "#ef4444", 5)

    if is_ko:
        note = "검증 관점: ECU6 구현은 소스, DBC/CAPL, CMM, host unit test, Debug_FLASH 산출물로 추적된다."
    else:
        note = "Verification view: ECU6 evidence is traced through source code, DBC/CAPL, CMM scripts, host unit tests, and Debug_FLASH artifacts."
    draw.rounded_rectangle((110, 1040, 2150, 1220), radius=24, fill="#ffffff", outline="#cbd5e1", width=3)
    centered_text(draw, (130, 1055, 2130, 1205), note, body_font, "#334155")
    img.save(path)
    return path


def make_static_structure(lang: str):
    is_ko = lang == "ko"
    path = OUT_DIR / f"compliance_static_structure_{lang}.png"
    img = Image.new("RGB", (2200, 1350), "#ffffff")
    draw = ImageDraw.Draw(img)
    title_font = font(FONT_KR_BOLD if is_ko else FONT_EN_BOLD, 54, FONT_EN_BOLD)
    body_font = font(FONT_KR if is_ko else FONT_EN, 28, FONT_EN)
    small_font = font(FONT_KR if is_ko else FONT_EN, 23, FONT_EN)
    title = "정적 구조 다이어그램: 파일/모듈 증거" if is_ko else "Static Structure Diagram: File and Module Evidence"
    label(draw, (70, 50), title, title_font, "#0b2545")
    cols = [
        ("Generated / SDK PAL\nGenerated_Code/*.c\nclock, CAN, UART, pin mux\nFreeRTOSConfig.h", (90, 210, 500, 1040), "#e0f2fe", "#0284c7"),
        ("Drivers\nSources/drivers/can.c\ngpio.c, uart.c\nCAN mailboxes, Tx/Rx tasks", (560, 210, 970, 1040), "#dcfce7", "#16a34a"),
        ("Services\nsystem_init.c\nlogger.c\nms_scheduler.c\ncom_gateway.c", (1030, 210, 1440, 1040), "#fef3c7", "#d97706"),
        ("Tests / Evidence\nTests/host/*.py\nCANOE/*.can, *.dbc\ncmm/*.cmm\nDebug_FLASH/*.elf, *.map", (1500, 210, 2110, 1040), "#f1f5f9", "#64748b"),
    ]
    for text, box, fill, outline in cols:
        rounded(draw, box, text, body_font, fill=fill, outline=outline)
    for start_x, end_x in [(500, 560), (970, 1030), (1440, 1500)]:
        arrow(draw, (start_x, 620), (end_x, 620), "#334155", 5)
    reqs = "REQ-COM-001~015\nREQ-ECU6-001~016\nREQ-INT-001/007/008/009"
    rounded(draw, (560, 1110, 1440, 1260), reqs, small_font, fill="#eef2ff", outline="#6366f1")
    arrow(draw, (1220, 1040), (1220, 1110), "#6366f1", 5)
    img.save(path)
    return path


def make_dynamic_sequence(lang: str):
    is_ko = lang == "ko"
    path = OUT_DIR / f"compliance_dynamic_sequence_{lang}.png"
    img = Image.new("RGB", (2200, 1350), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = font(FONT_KR_BOLD if is_ko else FONT_EN_BOLD, 52, FONT_EN_BOLD)
    body_font = font(FONT_KR if is_ko else FONT_EN, 25, FONT_EN)
    small_font = font(FONT_KR if is_ko else FONT_EN, 22, FONT_EN)
    title = "동적 다이어그램: CAN 수신 및 GatewayStatus 송신 흐름" if is_ko else "Dynamic Diagram: CAN Reception and GatewayStatus Transmission"
    label(draw, (70, 45), title, title_font, "#0b2545")
    actors = [
        ("CANoe / ECU1~5", 190),
        ("can.c\nvCANTask", 520),
        ("COM Gateway\ncom_gateway.c", 930),
        ("OTA/Security\nstate machine", 1370),
        ("CAN PAL\nTx/Rx", 1760),
        ("logger.c\nLOG_Print", 2040),
    ]
    for text, x in actors:
        rounded(draw, (x - 130, 145, x + 130, 255), text, small_font, fill="#dbeafe", outline="#93c5fd")
        draw.line((x, 255, x, 1235), fill="#94a3b8", width=3)
    steps = [
        (190, 520, "0x101~0x501 monitored frame" if not is_ko else "0x101~0x501 감시 frame"),
        (520, 930, "COM_GatewayHandleRxMessage()" if not is_ko else "COM_GatewayHandleRxMessage() 호출"),
        (930, 930, "Allowed ID check" if not is_ko else "허용 ID 확인"),
        (930, 930, "XOR checksum validation" if not is_ko else "XOR checksum 검증"),
        (930, 930, "Alive/timeout supervision" if not is_ko else "alive/timeout 감시"),
        (930, 2040, "Fault/recovery log" if not is_ko else "fault/recovery 로그"),
        (190, 1370, "0x650 OTA_Request" if not is_ko else "0x650 OTA_Request"),
        (1370, 1370, "Seed/key, AuthByte, checksum, lockout" if not is_ko else "seed/key, AuthByte, checksum, lockout"),
        (520, 930, "COM_GatewayBuildStatusMessage()" if not is_ko else "COM_GatewayBuildStatusMessage()"),
        (930, 520, "0x601 payload + alive + checksum" if not is_ko else "0x601 payload + alive + checksum"),
        (520, 1760, "CAN_SendMessage() every 100 ms" if not is_ko else "100 ms마다 CAN_SendMessage()"),
        (1760, 190, "GatewayStatus 0x601" if not is_ko else "GatewayStatus 0x601 송신"),
    ]
    y = 330
    for x0, x1, text in steps:
        if x0 == x1:
            rounded(draw, (x0 + 35, y - 26, x0 + 360, y + 50), text, small_font, fill="#ecfeff", outline="#67e8f9", width=2)
        else:
            arrow(draw, (x0 + (120 if x1 > x0 else -120), y), (x1 - (120 if x1 > x0 else -120), y), "#2563eb", 4)
            label(draw, (min(x0, x1) + 135, y - 36), text, small_font, "#1e293b")
        y += 80
    img.save(path)
    return path


def make_diagrams(lang: str):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    return [
        make_system_architecture(lang),
        make_static_structure(lang),
        make_dynamic_sequence(lang),
    ]


def compliance_summary(reqs, statuses, lang):
    counts = Counter(status_for(r["id"], statuses, lang) for r in reqs)
    if lang == "ko":
        order = ["충족", "부분 충족", "미충족", "확인 필요"]
        meaning = {
            "충족": "코드, 문서, 테스트 또는 실행 가능한 증거가 요구사항을 닫기에 충분함",
            "부분 충족": "구현 또는 테스트 자산은 있으나 런타임/외부 도구 증거가 부족함",
            "미충족": "의미 있는 구현 증거가 없음",
            "확인 필요": "현재 워크스페이스만으로는 판정할 증거가 부족함",
        }
        return [(k, counts[k], meaning[k]) for k in order]
    order = ["Compliant", "Partially Compliant", "Non-compliant", "Needs Confirmation"]
    meaning = {
        "Compliant": "Code, documentation, test, or executable evidence is sufficient to close the requirement.",
        "Partially Compliant": "Implementation or test assets exist, but runtime or external-tool evidence is incomplete.",
        "Non-compliant": "No meaningful implementation evidence was found.",
        "Needs Confirmation": "The current workspace does not contain enough evidence to make a final determination.",
    }
    return [(k, counts[k], meaning[k]) for k in order]


def matrix_rows(reqs, statuses, lang):
    rows = []
    for r in reqs:
        ev = get_evidence(r["id"])
        status = status_for(r["id"], statuses, lang)
        desc = ko_description(r["id"], r["statement"]) if lang == "ko" else r["statement"]
        if lang == "ko":
            note = ev["note"]
            if status == "확인 필요":
                note = "증거 부족 / 확인 필요. " + note
            rows.append([r["id"], desc, status, ev["impl"], ev["verify"], ev["evidence"], note])
        else:
            note = ev["note"]
            if status == "Needs Confirmation":
                note = "Insufficient evidence / needs confirmation. " + note
            rows.append([r["id"], desc, status, ev["impl"], ev["verify"], ev["evidence"], note])
    return rows


def add_common_sections(doc, reqs, statuses, lang):
    if lang == "ko":
        add_para(doc, "문서 목적", style="Heading 1")
        add_para(doc, "본 문서는 Automotive_Cluster 워크스페이스가 원 요구사항 문서의 요구사항을 어느 수준까지 충족하는지, 요구사항 ID별로 구현 위치, 검증 방법, 증거 파일/로그, 판정을 연결하여 제시한다. 증거가 부족한 항목은 충족으로 판정하지 않고 확인 필요로 분류하였다.")
        add_para(doc, "프로젝트 개요", style="Heading 1")
        add_para(doc, "원 요구사항 문서는 MPC5748G 기반 6개 ECU 교육 프로젝트를 정의한다. 현재 워크스페이스는 ECU6 Gateway + OTA + Security 노드를 중심으로 구성되어 있으며, ECU1~ECU5는 CANoe 시뮬레이터/DBC 지원 자산은 있으나 독립 MPC5748G ECU 구현 증거는 확인되지 않았다.")
        add_para(doc, "검증 범위", style="Heading 1")
        add_bullets(doc, [
            "전체 요구사항 96개를 추적성 매트릭스에 포함하였다.",
            "ECU6 활성 범위는 Common 15개, ECU6 16개, ECU6 관련 Integration 4개로 총 35개이다.",
            "ECU1~ECU5 개별 보드 요구사항과 일부 통합 시나리오는 현재 워크스페이스 증거만으로 최종 충족 판정하지 않았다.",
        ])
        add_para(doc, "사용한 자료 및 증거 소스", style="Heading 1")
        source_rows = [
            ("요구사항", "MPC5748G_Embedded_Training_Requirements_EN.docx", "96개 요구사항 추출"),
            ("상태 문서", "Documentation/ecu6_gateway/requirements_status.md", "ECU6 활성 범위 35개 판정 기준"),
            ("설계/시험 문서", "Documentation/ecu6_gateway/design.md; test_report.md", "CAN, GatewayStatus, OTA/security 설계 및 테스트 자산"),
            ("소스 코드", "Sources/**/*.c; Generated_Code/**/*.c", "초기화, CAN, gateway, scheduler, logging 구현"),
            ("CANoe/CMM", "CANOE/*; cmm/*", "DBC, CAPL, panel, Trace32 스크립트 및 binary report artifact"),
            ("빌드/테스트", "Debug_FLASH/Automotive_Cluster.elf; host_unit_tests_2026-06-04.txt", "빌드 산출물 및 9/9 host unit test 통과"),
        ]
        add_table(doc, ["분류", "경로", "검토 내용"], source_rows, [1700, 3800, 3860], lang, 9)
        add_para(doc, "요구사항 충족 요약", style="Heading 1")
        add_table(doc, ["판정", "수량", "의미"], compliance_summary(reqs, statuses, lang), [1700, 900, 6760], lang, 9)
    else:
        add_para(doc, "Document Purpose", style="Heading 1")
        add_para(doc, "This report states the compliance evidence for the Automotive_Cluster workspace by linking each requirement ID to its implementation location, verification method, evidence files/logs, and final status. Items without sufficient evidence are not marked compliant.")
        add_para(doc, "Project Overview", style="Heading 1")
        add_para(doc, "The source requirement document defines a six-ECU MPC5748G training project. This workspace is ECU6 Gateway + OTA + Security focused. CANoe simulator and DBC assets exist for ECU1 through ECU5, but independent MPC5748G ECU implementation evidence for those roles was not found.")
        add_para(doc, "Verification Scope", style="Heading 1")
        add_bullets(doc, [
            "All 96 source requirements are included in the traceability matrix.",
            "The ECU6 active scope contains 35 requirements: 15 Common, 16 ECU6, and 4 ECU6-related Integration requirements.",
            "ECU1 through ECU5 board-specific requirements and several integration scenarios remain Needs Confirmation based on the current workspace evidence.",
        ])
        add_para(doc, "Evidence Sources Reviewed", style="Heading 1")
        source_rows = [
            ("Requirements", "MPC5748G_Embedded_Training_Requirements_EN.docx", "Extracted 96 requirements"),
            ("Status document", "Documentation/ecu6_gateway/requirements_status.md", "Status basis for the 35 ECU6 active-scope requirements"),
            ("Design/test docs", "Documentation/ecu6_gateway/design.md; test_report.md", "CAN, GatewayStatus, OTA/security design and test assets"),
            ("Source code", "Sources/**/*.c; Generated_Code/**/*.c", "Initialization, CAN, gateway, scheduler, and logging implementation"),
            ("CANoe/CMM", "CANOE/*; cmm/*", "DBC, CAPL, panel, Trace32 scripts, and binary report artifacts"),
            ("Build/test", "Debug_FLASH/Automotive_Cluster.elf; host_unit_tests_2026-06-04.txt", "Build artifacts and 9/9 passing host unit tests"),
        ]
        add_table(doc, ["Category", "Path", "Reviewed Evidence"], source_rows, [1700, 3800, 3860], lang, 9)
        add_para(doc, "Compliance Summary", style="Heading 1")
        add_table(doc, ["Status", "Count", "Meaning"], compliance_summary(reqs, statuses, lang), [1700, 900, 6760], lang, 9)


def add_matrix_section(doc, reqs, statuses, lang):
    doc.add_section(WD_SECTION.NEW_PAGE)
    section = doc.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    title = "요구사항 추적성 매트릭스" if lang == "ko" else "Requirements Traceability Matrix"
    add_para(doc, title, style="Heading 1")
    headers = (
        ["요구사항 ID", "요구사항 내용", "충족 여부", "구현 위치", "검증 방법", "증거 파일/로그", "비고"]
        if lang == "ko"
        else ["Requirement ID", "Requirement Description", "Compliance Status", "Implementation Location", "Verification Method", "Evidence File/Log", "Notes"]
    )
    widths = [950, 2750, 1050, 2500, 2200, 2300, 2200]
    add_table(doc, headers, matrix_rows(reqs, statuses, lang), widths, lang, 6.7)


def add_diagram_block(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.35))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string("555555")
    r.italic = True


def add_detail_sections(doc, lang, diagrams):
    # Return to portrait.
    doc.add_section(WD_SECTION.NEW_PAGE)
    section = doc.sections[-1]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    if lang == "ko":
        add_para(doc, "기능별 상세 증빙", style="Heading 1")
        add_diagram_block(doc, diagrams[0], "그림 1. 시스템 아키텍처: ECU6, CANoe, Trace32, S32DS 산출물의 증거 연결")
        add_diagram_block(doc, diagrams[1], "그림 2. 정적 구조: Generated/SDK, Driver, Service, Test/Evidence 계층")
        add_diagram_block(doc, diagrams[2], "그림 3. 동적 흐름: CAN 수신, gateway 검증, OTA/security, 0x601 송신")
        add_table(doc, ["기능 영역", "구현 위치", "주요 증거"], [
            ("초기화/RTOS", "Sources/main.c; Sources/services/system_init.c; Generated_Code/FreeRTOSConfig.h", "clock/pin/UART/CAN 초기화, FreeRTOS task 생성, 32 KB heap 설정"),
            ("CAN gateway", "Sources/drivers/can.c; Sources/services/com_gateway.c", "0x101~0x501 수신, 0x601 송신, unknown ID catch-all, timeout/checksum/alive 감시"),
            ("OTA/security", "Sources/services/com_gateway.c; CANOE/ECU6_OTA_State_Test.can", "0x650/0x651, START/DATA/END/ACTIVATE/ABORT, seed/key, AuthByte, lockout, ActiveSWVersion"),
            ("Trace32", "cmm/init_debug.cmm; cmm/test_debug.cmm; cmm/gateway_debug.cmm; cmm/ota_debug.cmm; cmm/security_debug.cmm", "디버그 진입, breakpoint, watch 변수, interrupt/scheduler evidence capture script"),
            ("CANoe", "CANOE/6_ECUs_dbc_0206.dbc; CANOE/*.can; Configuration1.cfg", "6 ECU DBC, ECU simulator CAPL, Gateway/OTA CAPL test, panel bridge 및 binary report artifact"),
        ], [1900, 3500, 3960], lang, 8.5)
        add_para(doc, "빌드 및 실행 검증 결과", style="Heading 1")
        add_bullets(doc, [
            "Debug_FLASH/Automotive_Cluster.elf(2,342,192 bytes)와 Debug_FLASH/Automotive_Cluster.map(594,806 bytes)이 확인되었다.",
            "map 파일에서 vCANTask, OTA_HandleRequestMessage, COM_GatewayBuildStatusMessage, vComSchedulerTask 심볼이 확인되었다.",
            "현재 셸 PATH에서 powerpc-eabivle-gcc 및 make는 확인되지 않아 Codex가 S32DS rebuild를 수행하지 않았다.",
            "S32DS flash/run 10분 안정성, UART boot log, Trace32 screenshot, CANoe readable verdict export는 아직 확인 필요이다.",
        ])
        add_para(doc, "테스트 검증 결과", style="Heading 1")
        add_table(doc, ["테스트", "명령/도구", "결과", "관련 요구사항"], [
            ("Host unit tests", "python -m unittest discover -s Tests/host -p \"test_*.py\" -v", "9/9 pass", "REQ-COM-005/007/015, REQ-ECU6-002/003/004/009/010/011/013/015, REQ-INT-007/008"),
            ("Gateway CAPL tests", "CANOE/ECU6_Gateway_Test.can", "테스트 자산 및 binary report 존재, readable verdict 필요", "REQ-COM-003/011, REQ-ECU6-001/002/003/006/014, REQ-INT-009"),
            ("OTA/security CAPL tests", "CANOE/ECU6_OTA_State_Test.can", "테스트 자산 및 binary report 존재, readable verdict 필요", "REQ-ECU6-005~013/015, REQ-INT-007/008"),
        ], [1800, 3300, 2500, 1760], lang, 8.2)
        add_para(doc, "미확인/증거 부족 항목", style="Heading 1")
        add_bullets(doc, [
            "REQ-COM-001, REQ-COM-003, REQ-COM-004, REQ-INT-001은 구현/자산 증거는 있으나 최종 런타임 증거가 부족하여 부분 충족으로 판정하였다.",
            "ECU1~ECU5 개별 요구사항 56개와 REQ-INT-002~006은 현재 워크스페이스의 ECU6 중심 증거만으로 충족 판정할 수 없다.",
            "CANoe .vtestreport 파일은 존재하지만 텍스트 검색으로 PASS/FAIL verdict를 확인할 수 없어 최종 판정 증거로 사용하지 않았다.",
        ])
        add_para(doc, "결론", style="Heading 1")
        add_para(doc, "Automotive_Cluster 워크스페이스는 ECU6 Gateway + OTA + Security 활성 범위의 대부분 요구사항을 코드/문서/host unit test 수준에서 충족한다. 다만 전체 96개 요구사항을 모두 충족한다고 증명하기에는 ECU1~ECU5 보드 구현 증거와 일부 통합 실행 증거가 부족하다. 따라서 본 보고서는 31개 충족, 4개 부분 충족, 61개 확인 필요로 판정한다.")
        add_para(doc, "부록: 파일 경로, 로그 요약, 테스트 명령, 주요 설정값", style="Heading 1")
        add_table(doc, ["항목", "값"], [
            ("요구사항 원문", str(REQ_DOC.relative_to(ROOT))),
            ("Host test log", str(HOST_TEST_LOG.relative_to(ROOT))),
            ("Host test command", r"C:\Users\JunsangS\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe -m unittest discover -s Tests/host -p \"test_*.py\" -v"),
            ("CAN ID set", "0x101, 0x201, 0x301, 0x401, 0x501, 0x601, 0x650, 0x651"),
            ("핵심 설정", "CAN_NODE_COUNT=6, CAN_TX_PERIOD_MS=100, LOG_BUFFER_SIZE=256, configTOTAL_HEAP_SIZE=32768"),
        ], [2300, 7060], lang, 8.5)
    else:
        add_para(doc, "Detailed Evidence by Function", style="Heading 1")
        add_diagram_block(doc, diagrams[0], "Figure 1. System architecture: evidence linkage across ECU6, CANoe, Trace32, and S32DS artifacts")
        add_diagram_block(doc, diagrams[1], "Figure 2. Static structure: Generated/SDK, Driver, Service, and Test/Evidence layers")
        add_diagram_block(doc, diagrams[2], "Figure 3. Dynamic flow: CAN reception, gateway validation, OTA/security, and 0x601 transmission")
        add_table(doc, ["Function Area", "Implementation Location", "Key Evidence"], [
            ("Initialization/RTOS", "Sources/main.c; Sources/services/system_init.c; Generated_Code/FreeRTOSConfig.h", "Clock/pin/UART/CAN initialization, FreeRTOS task creation, 32 KB heap setting"),
            ("CAN gateway", "Sources/drivers/can.c; Sources/services/com_gateway.c", "0x101~0x501 reception, 0x601 transmission, unknown-ID catch-all, timeout/checksum/alive supervision"),
            ("OTA/security", "Sources/services/com_gateway.c; CANOE/ECU6_OTA_State_Test.can", "0x650/0x651, START/DATA/END/ACTIVATE/ABORT, seed/key, AuthByte, lockout, ActiveSWVersion"),
            ("Trace32", "cmm/init_debug.cmm; cmm/test_debug.cmm; cmm/gateway_debug.cmm; cmm/ota_debug.cmm; cmm/security_debug.cmm", "Debug entry, breakpoints, watched variables, interrupt/scheduler evidence capture scripts"),
            ("CANoe", "CANOE/6_ECUs_dbc_0206.dbc; CANOE/*.can; Configuration1.cfg", "Six-ECU DBC, ECU simulator CAPL, Gateway/OTA CAPL tests, panel bridge, and binary report artifacts"),
        ], [1900, 3500, 3960], lang, 8.5)
        add_para(doc, "Build and Execution Verification Results", style="Heading 1")
        add_bullets(doc, [
            "Debug_FLASH/Automotive_Cluster.elf (2,342,192 bytes) and Debug_FLASH/Automotive_Cluster.map (594,806 bytes) were observed.",
            "The map file contains vCANTask, OTA_HandleRequestMessage, COM_GatewayBuildStatusMessage, and vComSchedulerTask symbols.",
            "powerpc-eabivle-gcc and make were not found in the current shell PATH, so Codex did not perform an S32DS rebuild.",
            "S32DS flash/run stability for 10 minutes, UART boot logs, Trace32 screenshots, and readable CANoe verdict exports still require confirmation.",
        ])
        add_para(doc, "Test Verification Results", style="Heading 1")
        add_table(doc, ["Test", "Command/Tool", "Result", "Related Requirements"], [
            ("Host unit tests", "python -m unittest discover -s Tests/host -p \"test_*.py\" -v", "9/9 pass", "REQ-COM-005/007/015, REQ-ECU6-002/003/004/009/010/011/013/015, REQ-INT-007/008"),
            ("Gateway CAPL tests", "CANOE/ECU6_Gateway_Test.can", "Test assets and binary report exist; readable verdict export required", "REQ-COM-003/011, REQ-ECU6-001/002/003/006/014, REQ-INT-009"),
            ("OTA/security CAPL tests", "CANOE/ECU6_OTA_State_Test.can", "Test assets and binary report exist; readable verdict export required", "REQ-ECU6-005~013/015, REQ-INT-007/008"),
        ], [1800, 3300, 2500, 1760], lang, 8.2)
        add_para(doc, "Unconfirmed or Insufficient Evidence Items", style="Heading 1")
        add_bullets(doc, [
            "REQ-COM-001, REQ-COM-003, REQ-COM-004, and REQ-INT-001 are Partially Compliant because implementation/test assets exist but final runtime evidence is incomplete.",
            "The 56 ECU1 through ECU5 board-specific requirements and REQ-INT-002 through REQ-INT-006 cannot be marked compliant from the ECU6-focused workspace evidence.",
            "CANoe .vtestreport files exist, but PASS/FAIL verdicts were not readable by text search and were therefore not used as final verdict evidence.",
        ])
        add_para(doc, "Conclusion", style="Heading 1")
        add_para(doc, "The Automotive_Cluster workspace satisfies most of the ECU6 Gateway + OTA + Security active scope at code, documentation, and host-unit-test level. It does not yet provide enough evidence to prove full compliance for all 96 requirements because ECU1 through ECU5 board implementations and several full-system runtime artifacts are missing. This report therefore records 31 Compliant, 4 Partially Compliant, and 61 Needs Confirmation requirements.")
        add_para(doc, "Appendix: File Paths, Log Summary, Test Commands, Key Configuration Values", style="Heading 1")
        add_table(doc, ["Item", "Value"], [
            ("Source requirement document", str(REQ_DOC.relative_to(ROOT))),
            ("Host test log", str(HOST_TEST_LOG.relative_to(ROOT))),
            ("Host test command", r"C:\Users\JunsangS\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe -m unittest discover -s Tests/host -p \"test_*.py\" -v"),
            ("CAN ID set", "0x101, 0x201, 0x301, 0x401, 0x501, 0x601, 0x650, 0x651"),
            ("Key settings", "CAN_NODE_COUNT=6, CAN_TX_PERIOD_MS=100, LOG_BUFFER_SIZE=256, configTOTAL_HEAP_SIZE=32768"),
        ], [2300, 7060], lang, 8.5)


def add_headers_footers(doc, lang):
    label = "Automotive_Cluster Requirements Compliance" if lang == "en" else "Automotive_Cluster 요구사항 충족 증명"
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = label
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor.from_string("555555")
        footer = section.footer.paragraphs[0]
        footer.text = "Generated evidence report"
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.runs[0].font.size = Pt(8)
        footer.runs[0].font.color.rgb = RGBColor.from_string("777777")


def build(lang: str, out_name: str):
    reqs = load_requirements()
    statuses = load_statuses()
    diagrams = make_diagrams(lang)
    doc = Document()
    set_doc_styles(doc, lang)
    add_cover(doc, lang)
    add_common_sections(doc, reqs, statuses, lang)
    add_matrix_section(doc, reqs, statuses, lang)
    add_detail_sections(doc, lang, diagrams)
    add_headers_footers(doc, lang)
    out = ROOT / out_name
    doc.core_properties.title = (
        "Automotive_Cluster 프로젝트 요구사항 충족 증명 보고서"
        if lang == "ko"
        else "Automotive_Cluster Requirements Compliance Evidence Report"
    )
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.subject = "Requirements compliance evidence"
    doc.save(out)
    return out


if __name__ == "__main__":
    ko = build("ko", "Automotive_Cluster_Requirements_Compliance_Report_KO.docx")
    en = build("en", "Automotive_Cluster_Requirements_Compliance_Report_EN.docx")
    print(ko)
    print(en)
