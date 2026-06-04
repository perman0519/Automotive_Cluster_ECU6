from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "Documentation" / "generated"
FONT_KR = Path("C:/Windows/Fonts/malgun.ttf")
FONT_KR_BOLD = Path("C:/Windows/Fonts/malgunbd.ttf")
FONT_MONO = Path("C:/Windows/Fonts/consola.ttf")


REQ_ROWS_KO = [
    ("REQ-COM-005", "P0", "Critical CAN/external input timeout monitoring", "0x101~0x501 수신 중단을 300 ms 기준으로 감시하고, timeout 시 AliveError/GatewayState/log에 반영"),
    ("REQ-COM-006", "P0", "4-bit alive counter in primary status message", "0x601 Byte6 lower nibble에 aliveCounter를 넣고 0~15 wrap"),
    ("REQ-COM-007", "P0", "Simple checksum in primary status message", "0x601 송신 checksum 생성 및 0x101~0x501 수신 checksum 검증"),
    ("REQ-ECU6-001", "P0", "Monitor or route allowed CAN IDs", "0x101, 0x201, 0x301, 0x401, 0x501만 monitored ID로 허용"),
    ("REQ-ECU6-002", "P0", "Block or count unknown CAN IDs", "허용되지 않은 ID 수신 시 BlockedMsgCnt 증가 및 log 출력"),
    ("REQ-ECU6-003", "P0", "Validate checksum for monitored primary messages", "checksum mismatch 시 ChecksumError set, BlockedMsgCnt 증가, log 출력"),
    ("REQ-ECU6-004", "P0", "Monitor alive counters for selected ECUs", "동일 alive counter가 3회 반복되면 AliveError set"),
    ("REQ-ECU6-014", "P0", "Send GatewayStatus on CAN ID 0x601 every 100 ms", "vCANTask가 100 ms마다 COM_GatewayBuildStatusMessage() 후 CAN_SendMessage() 호출"),
]

REQ_ROWS_EN = [
    ("REQ-COM-005", "P0", "Critical CAN/external input timeout monitoring", "Monitors 0x101~0x501 reception with a 300 ms timeout and reports faults through AliveError, GatewayState, and logs"),
    ("REQ-COM-006", "P0", "4-bit alive counter in primary status message", "Places aliveCounter in the lower nibble of 0x601 Byte6 and wraps 0~15"),
    ("REQ-COM-007", "P0", "Simple checksum in primary status message", "Generates checksum for 0x601 and validates received checksums for 0x101~0x501"),
    ("REQ-ECU6-001", "P0", "Monitor or route allowed CAN IDs", "Accepts only 0x101, 0x201, 0x301, 0x401, and 0x501 as monitored IDs"),
    ("REQ-ECU6-002", "P0", "Block or count unknown CAN IDs", "Increments BlockedMsgCnt and logs when an unknown ID is received"),
    ("REQ-ECU6-003", "P0", "Validate checksum for monitored primary messages", "Sets ChecksumError, increments BlockedMsgCnt, and logs on checksum mismatch"),
    ("REQ-ECU6-004", "P0", "Monitor alive counters for selected ECUs", "Sets AliveError when the same alive counter is repeated 3 times"),
    ("REQ-ECU6-014", "P0", "Send GatewayStatus on CAN ID 0x601 every 100 ms", "vCANTask calls COM_GatewayBuildStatusMessage() and CAN_SendMessage() every 100 ms"),
]


FUNCTION_ROWS_KO = [
    ("can.c", "CAN_ConfigGatewayMailboxes()", "ECU6 Gateway용 RX/TX mailbox를 설정한다."),
    ("can.c", "CAN_ReceiveMessage()", "CAN PAL의 CAN_ReceiveBlocking()을 감싼 수신 wrapper이다."),
    ("can.c", "CAN_SendMessage()", "완성된 CAN frame을 ECU6 TX mailbox로 송신한다."),
    ("can.c", "vCANTask()", "수신 polling, timeout supervision, 100 ms GatewayStatus 송신을 주기적으로 수행한다."),
    ("com_gateway.c", "COM_GatewayInit()", "monitor table, GatewayStatus 상태값, alive/checksum/error flag를 초기화한다."),
    ("com_gateway.c", "COM_GatewayHandleRxMessage()", "수신 frame의 ID 허용 여부, checksum, alive counter, timeout recovery를 처리한다."),
    ("com_gateway.c", "COM_GatewayCheckTimeouts()", "각 monitored ID의 lastRxTick을 기준으로 300 ms timeout을 판정한다."),
    ("com_gateway.c", "COM_GatewayBuildStatusMessage()", "0x601 GatewayStatus payload, alive counter, checksum을 생성한다."),
    ("com_gateway.c", "COM_CalculateChecksum()", "CAN_ID_LSB와 data[0]~data[6]을 XOR하여 checksum을 계산한다."),
    ("com_gateway.c", "COM_VerifyChecksum()", "수신 data[7]과 계산 checksum을 비교한다."),
    ("com_gateway.c", "COM_UpdateAliveMonitor()", "Byte6 lower nibble alive counter가 멈췄는지 감시한다."),
    ("automotive.h", "function prototypes", "CAN wrapper와 COM Gateway service 함수 시그니처를 공개한다."),
]

FUNCTION_ROWS_EN = [
    ("can.c", "CAN_ConfigGatewayMailboxes()", "Configures RX/TX mailboxes used by ECU6 gateway communication."),
    ("can.c", "CAN_ReceiveMessage()", "Wraps CAN PAL CAN_ReceiveBlocking() for gateway reception."),
    ("can.c", "CAN_SendMessage()", "Transmits a prepared CAN frame through the ECU6 TX mailbox."),
    ("can.c", "vCANTask()", "Runs receive polling, timeout supervision, and 100 ms GatewayStatus transmission."),
    ("com_gateway.c", "COM_GatewayInit()", "Initializes the monitor table, GatewayStatus state, alive/checksum/error flags."),
    ("com_gateway.c", "COM_GatewayHandleRxMessage()", "Processes received frame ID filtering, checksum, alive counter, and timeout recovery."),
    ("com_gateway.c", "COM_GatewayCheckTimeouts()", "Checks 300 ms timeout using each monitored ID's lastRxTick."),
    ("com_gateway.c", "COM_GatewayBuildStatusMessage()", "Builds 0x601 GatewayStatus payload, alive counter, and checksum."),
    ("com_gateway.c", "COM_CalculateChecksum()", "Calculates checksum by XORing CAN_ID_LSB with data[0]~data[6]."),
    ("com_gateway.c", "COM_VerifyChecksum()", "Compares received data[7] against the calculated checksum."),
    ("com_gateway.c", "COM_UpdateAliveMonitor()", "Supervises whether the Byte6 lower-nibble alive counter stops changing."),
    ("automotive.h", "function prototypes", "Publishes CAN wrapper and COM Gateway service signatures."),
]


PAYLOAD_ROWS_KO = [
    ("Byte0", "GatewayState", "0=INIT, 1=NORMAL, 2=DEGRADED"),
    ("Byte1", "OTA/security/error flags", "bit0~2 OTAState, bit3 SecurityUnlocked, bit4 SecurityLocked, bit5 AliveError, bit6 ChecksumError"),
    ("Byte2", "BlockedMsgCnt", "unknown ID 또는 checksum error 발생 시 saturation 증가"),
    ("Byte3", "SecurityErrorCounter", "보안 오류 카운터 자리, 현재 초기값 유지"),
    ("Byte4", "ActiveSWVersion", "현재 활성 SW version"),
    ("Byte5", "GatewayRxMask", "ECU1~ECU5 monitored message 수신 여부 bit mask"),
    ("Byte6", "AliveCounter", "lower 4 bit 사용, 0~15 wrap"),
    ("Byte7", "Checksum", "CAN_ID_LSB ^ Byte0 ^ ... ^ Byte6"),
]

PAYLOAD_ROWS_EN = [
    ("Byte0", "GatewayState", "0=INIT, 1=NORMAL, 2=DEGRADED"),
    ("Byte1", "OTA/security/error flags", "bit0~2 OTAState, bit3 SecurityUnlocked, bit4 SecurityLocked, bit5 AliveError, bit6 ChecksumError"),
    ("Byte2", "BlockedMsgCnt", "Saturating count for unknown ID or checksum error"),
    ("Byte3", "SecurityErrorCounter", "Reserved security error counter, currently initialized"),
    ("Byte4", "ActiveSWVersion", "Currently active software version"),
    ("Byte5", "GatewayRxMask", "Bit mask showing received monitored ECU1~ECU5 messages"),
    ("Byte6", "AliveCounter", "Lower 4 bits, wraps 0~15"),
    ("Byte7", "Checksum", "CAN_ID_LSB ^ Byte0 ^ ... ^ Byte6"),
]


def load_font(path, size):
    return ImageFont.truetype(str(path), size)


def draw_wrapped_center(draw, box, text, font, fill=(30, 41, 59), spacing=8):
    x0, y0, x1, y1 = box
    max_width = x1 - x0 - 28
    words = text.split()
    lines = []
    line = ""
    for word in words:
        candidate = word if not line else f"{line} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    heights = [draw.textbbox((0, 0), ln, font=font)[3] - draw.textbbox((0, 0), ln, font=font)[1] for ln in lines]
    total_h = sum(heights) + spacing * (len(lines) - 1)
    y = y0 + ((y1 - y0) - total_h) / 2
    for ln, h in zip(lines, heights):
        tw = draw.textbbox((0, 0), ln, font=font)[2]
        draw.text((x0 + ((x1 - x0) - tw) / 2, y), ln, font=font, fill=fill)
        y += h + spacing


def arrow(draw, start, end, fill=(51, 65, 85), width=4):
    draw.line([start, end], fill=fill, width=width)
    x0, y0 = start
    x1, y1 = end
    if x1 >= x0:
        points = [(x1, y1), (x1 - 18, y1 - 9), (x1 - 18, y1 + 9)]
    else:
        points = [(x1, y1), (x1 + 18, y1 - 9), (x1 + 18, y1 + 9)]
    draw.polygon(points, fill=fill)


def rounded_box(draw, box, text, font, fill, outline=(71, 85, 105), text_fill=(15, 23, 42)):
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=3)
    draw_wrapped_center(draw, box, text, font, fill=text_fill)


def make_sequence_diagram(lang):
    is_ko = lang == "ko"
    path = OUT_DIR / f"can_sequence_{lang}.png"
    img = Image.new("RGB", (2200, 1350), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = load_font(FONT_KR_BOLD, 54)
    font = load_font(FONT_KR, 30)
    small = load_font(FONT_KR, 24)
    header_fill = "#dbeafe"
    accent = "#2563eb"

    title = "ECU6 CAN Gateway 함수 시퀀스" if is_ko else "ECU6 CAN Gateway Function Sequence"
    draw.text((70, 45), title, font=title_font, fill="#0f172a")
    participants = [
        ("main.c\nmain()", 190),
        ("can.c\nvCANTask()", 520),
        ("can.c\nCAN wrapper", 850),
        ("com_gateway.c\nCOM Gateway", 1220),
        ("CAN PAL", 1600),
        ("logger.c\nLOG_Print()", 1930),
    ]
    for label, x in participants:
        rounded_box(draw, (x - 130, 140, x + 130, 240), label, small, header_fill, outline="#93c5fd")
        draw.line((x, 240, x, 1260), fill="#94a3b8", width=3)

    if is_ko:
        steps = [
            (190, 520, "xTaskCreate(vCANTask)"),
            (520, 850, "CAN_ConfigGatewayMailboxes()"),
            (520, 1220, "COM_GatewayInit()"),
            (520, 850, "CAN_ReceiveMessage()"),
            (850, 1600, "CAN_ReceiveBlocking()"),
            (520, 1220, "COM_GatewayHandleRxMessage(rxMsg)"),
            (1220, 1220, "ID 허용 여부 확인"),
            (1220, 1220, "checksum 검증"),
            (1220, 1220, "alive counter 감시"),
            (1220, 1930, "fault/recovery log"),
            (520, 1220, "COM_GatewayCheckTimeouts()"),
            (520, 1220, "COM_GatewayBuildStatusMessage(txMsg)"),
            (520, 850, "CAN_SendMessage(0x601)"),
            (850, 1600, "CAN_SendBlocking()"),
        ]
    else:
        steps = [
            (190, 520, "xTaskCreate(vCANTask)"),
            (520, 850, "CAN_ConfigGatewayMailboxes()"),
            (520, 1220, "COM_GatewayInit()"),
            (520, 850, "CAN_ReceiveMessage()"),
            (850, 1600, "CAN_ReceiveBlocking()"),
            (520, 1220, "COM_GatewayHandleRxMessage(rxMsg)"),
            (1220, 1220, "Check allowed ID"),
            (1220, 1220, "Verify checksum"),
            (1220, 1220, "Monitor alive counter"),
            (1220, 1930, "Fault/recovery log"),
            (520, 1220, "COM_GatewayCheckTimeouts()"),
            (520, 1220, "COM_GatewayBuildStatusMessage(txMsg)"),
            (520, 850, "CAN_SendMessage(0x601)"),
            (850, 1600, "CAN_SendBlocking()"),
        ]

    y = 310
    for x0, x1, label in steps:
        if x0 == x1:
            draw.rounded_rectangle((x0 + 30, y - 18, x0 + 310, y + 42), radius=14, fill="#ecfeff", outline="#67e8f9", width=2)
            draw_wrapped_center(draw, (x0 + 30, y - 18, x0 + 310, y + 42), label, small)
        else:
            arrow(draw, (x0 + 120, y), (x1 - 120, y), fill=accent, width=4)
            draw.text((min(x0, x1) + 130, y - 36), label, font=small, fill="#1e293b")
        y += 78

    note = "100 ms마다 0x601 GatewayStatus 송신" if is_ko else "0x601 GatewayStatus is transmitted every 100 ms"
    draw.rounded_rectangle((1320, 1170, 2100, 1260), radius=18, fill="#fef3c7", outline="#f59e0b", width=3)
    draw_wrapped_center(draw, (1320, 1170, 2100, 1260), note, font)
    img.save(path)
    return path


def make_architecture_diagram(lang):
    is_ko = lang == "ko"
    path = OUT_DIR / f"can_architecture_{lang}.png"
    img = Image.new("RGB", (2200, 1250), "#f8fafc")
    draw = ImageDraw.Draw(img)
    title_font = load_font(FONT_KR_BOLD, 54)
    font = load_font(FONT_KR, 30)
    small = load_font(FONT_KR, 25)

    title = "ECU6 CAN Gateway 아키텍처" if is_ko else "ECU6 CAN Gateway Architecture"
    draw.text((70, 45), title, font=title_font, fill="#0f172a")

    layers = [
        (80, 150, 2120, 300, "#e0f2fe", "Application / RTOS Layer", ["main.c main()", "can.c vCANTask()"]),
        (80, 350, 2120, 500, "#dcfce7", "Service Layer", ["com_gateway.c COM Gateway Service", "logger.c LOG_Print()"]),
        (80, 550, 2120, 700, "#fef3c7", "Driver Wrapper Layer", ["CAN_ConfigGatewayMailboxes()", "CAN_SendMessage()", "CAN_ReceiveMessage()"]),
        (80, 750, 2120, 900, "#ede9fe", "NXP PAL / Generated Layer", ["CAN PAL", "UART PAL"]),
        (80, 950, 2120, 1130, "#fee2e2", "External CAN Network", ["ECU1 0x101", "ECU2 0x201", "ECU3 0x301", "ECU4 0x401", "ECU5 0x501", "CANoe / All"]),
    ]
    if is_ko:
        layers[0] = (80, 150, 2120, 300, "#e0f2fe", "Application / RTOS 계층", ["main.c main()", "can.c vCANTask()"])
        layers[1] = (80, 350, 2120, 500, "#dcfce7", "Service 계층", ["com_gateway.c ECU6 Gateway Service", "logger.c LOG_Print()"])
        layers[2] = (80, 550, 2120, 700, "#fef3c7", "CAN Driver Wrapper 계층", ["CAN_ConfigGatewayMailboxes()", "CAN_SendMessage()", "CAN_ReceiveMessage()"])
        layers[3] = (80, 750, 2120, 900, "#ede9fe", "NXP PAL / Generated 계층", ["CAN PAL", "UART PAL"])
        layers[4] = (80, 950, 2120, 1130, "#fee2e2", "외부 CAN Network", ["ECU1 0x101", "ECU2 0x201", "ECU3 0x301", "ECU4 0x401", "ECU5 0x501", "CANoe / All"])

    for x0, y0, x1, y1, fill, label, boxes in layers:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=28, fill=fill, outline="#64748b", width=3)
        draw.text((x0 + 28, y0 + 18), label, font=font, fill="#0f172a")
        count = len(boxes)
        gap = 24
        box_w = (x1 - x0 - 310 - gap * (count - 1)) / count
        bx = x0 + 280
        for item in boxes:
            rounded_box(draw, (int(bx), y0 + 65, int(bx + box_w), y1 - 24), item, small, "#ffffff", outline="#94a3b8")
            bx += box_w + gap

    for y0, y1 in [(300, 350), (500, 550), (700, 750), (900, 950)]:
        arrow(draw, (1100, y0 + 10), (1100, y1 - 10), fill="#334155", width=5)

    img.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "1F4E79")
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=8.5)
        if widths:
            cell.width = Inches(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value), size=8.2)
            if widths:
                cells[idx].width = Inches(widths[idx])
    return table


def set_doc_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Malgun Gothic"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    for style_name, size, color in [
        ("Title", 24, "1F4E79"),
        ("Heading 1", 16, "1F4E79"),
        ("Heading 2", 12.5, "2F5597"),
        ("Heading 3", 11, "2F5597"),
    ]:
        style = styles[style_name]
        style.font.name = "Malgun Gothic"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(9.5)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    return p


def add_code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F3F6FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for line_i, line in enumerate(code.splitlines()):
        if line_i:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.2)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(9.2)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")


def create_doc(lang, seq_img, arch_img):
    is_ko = lang == "ko"
    doc = Document()
    set_doc_defaults(doc)

    if is_ko:
        title = "ECU6 CAN Gateway 요구사항 구현 설계서"
        subtitle = "REQ-COM-005/006/007 및 REQ-ECU6-001~004/014 구현 정리"
        out = OUT_DIR / "ECU6_CAN_Gateway_Requirements_KO.docx"
        req_rows, function_rows, payload_rows = REQ_ROWS_KO, FUNCTION_ROWS_KO, PAYLOAD_ROWS_KO
    else:
        title = "ECU6 CAN Gateway Requirements Implementation Design"
        subtitle = "Summary for REQ-COM-005/006/007 and REQ-ECU6-001~004/014"
        out = OUT_DIR / "ECU6_CAN_Gateway_Requirements_EN.docx"
        req_rows, function_rows, payload_rows = REQ_ROWS_EN, FUNCTION_ROWS_EN, PAYLOAD_ROWS_EN

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(23)
    r.font.color.rgb = RGBColor.from_string("1F4E79")
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("475569")
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

    add_para(doc, "Source files: Sources/drivers/can.c, Sources/services/com_gateway.c, Sources/automotive.h")
    add_para(doc, "Generated: 2026-05-22")

    add_heading(doc, "1. Overview" if not is_ko else "1. 개요")
    if is_ko:
        add_para(doc, "이번 구현은 CAN driver와 ECU6 Gateway service를 분리하여 CAN 하드웨어 접근과 요구사항 로직의 책임을 나누는 구조이다.")
        add_bullets(doc, [
            "can.c는 CAN PAL wrapper, mailbox 설정, 수신 polling, 100 ms 송신 trigger를 담당한다.",
            "com_gateway.c는 허용 CAN ID 감시, timeout, alive counter, checksum, GatewayStatus payload 생성을 담당한다.",
            "0x601 GatewayStatus는 ECU6가 All/CANoe로 송신하는 primary status message이다.",
        ])
    else:
        add_para(doc, "The implementation separates CAN driver responsibilities from ECU6 Gateway service logic.")
        add_bullets(doc, [
            "can.c handles CAN PAL wrapping, mailbox setup, receive polling, and the 100 ms transmit trigger.",
            "com_gateway.c handles allowed CAN ID monitoring, timeout, alive counter, checksum, and GatewayStatus payload creation.",
            "0x601 GatewayStatus is the ECU6 primary status message transmitted to All/CANoe.",
        ])

    add_heading(doc, "2. Requirement Mapping" if not is_ko else "2. 요구사항 매핑")
    headers = ["Requirement", "Priority", "Purpose", "Implementation"] if not is_ko else ["요구사항", "우선순위", "목적", "구현 내용"]
    add_table(doc, headers, req_rows, widths=[1.05, 0.65, 2.15, 3.05])

    add_heading(doc, "3. Software Architecture" if not is_ko else "3. 소프트웨어 아키텍처")
    add_para(doc, "Supervision is in com_gateway.c; CAN hardware access remains in can.c." if not is_ko else "아키텍처는 요구사항 감시 로직을 com_gateway.c에 집중시키고, CAN 하드웨어 접근은 can.c에 남기는 구조이다.")
    doc.add_picture(str(arch_img), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "4. Function Sequence" if not is_ko else "4. 함수 시퀀스")
    add_para(doc, "The sequence below shows how a received frame is validated and how 0x601 is periodically transmitted." if not is_ko else "아래 시퀀스는 수신 frame 검증과 0x601 주기 송신이 어떤 함수 흐름으로 처리되는지 보여준다.")
    doc.add_picture(str(seq_img), width=Inches(6.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "5. Function Responsibility Table" if not is_ko else "5. 함수 책임 표")
    headers = ["File", "Function", "Responsibility"] if not is_ko else ["파일", "함수", "역할"]
    add_table(doc, headers, function_rows, widths=[1.35, 2.2, 3.35])

    doc.add_page_break()
    add_heading(doc, "6. GatewayStatus 0x601 Payload" if not is_ko else "6. GatewayStatus 0x601 Payload")
    add_para(doc, "The current DLC 8 layout packs error flags into Byte1 and keeps Byte6/Byte7 for alive/checksum." if not is_ko else "현재 DLC 8 layout은 error flag를 Byte1에 pack하고, Byte6/Byte7을 alive/checksum에 사용한다.")
    headers = ["Byte", "Signal", "Meaning"] if not is_ko else ["Byte", "Signal", "의미"]
    add_table(doc, headers, payload_rows, widths=[0.75, 2.1, 4.05])

    add_heading(doc, "7. Checksum and Alive Counter Rules" if not is_ko else "7. Checksum 및 Alive Counter 규칙")
    checksum_text = (
        "Checksum is calculated with the requested simple XOR rule. The same function is used for GatewayStatus generation and monitored primary-message validation."
        if not is_ko
        else "Checksum은 요청한 simple XOR 규칙으로 계산한다. 동일 계산 함수가 GatewayStatus 송신과 monitored primary message 수신 검증에 사용된다."
    )
    add_para(doc, checksum_text)
    add_code_block(doc, """checksum = CAN_ID_LSB ^ data[0] ^ data[1] ^ data[2]
         ^ data[3] ^ data[4] ^ data[5] ^ data[6];""")
    if is_ko:
        add_bullets(doc, [
            "AliveCounter는 Byte6 lower 4 bit를 사용한다.",
            "ECU6 송신 aliveCounter는 0~15까지 증가 후 0으로 wrap한다.",
            "수신 alive counter가 3회 연속 동일하면 alive freeze fault로 판단한다.",
            "Checksum mismatch는 GatewayChecksumError를 set하고 BlockedMsgCnt를 증가시킨다.",
        ])
    else:
        add_bullets(doc, [
            "AliveCounter uses the lower 4 bits of Byte6.",
            "ECU6 transmit aliveCounter increments from 0 to 15 and wraps to 0.",
            "A received alive counter repeated for 3 consecutive checks is treated as an alive freeze fault.",
            "Checksum mismatch sets GatewayChecksumError and increments BlockedMsgCnt.",
        ])

    add_heading(doc, "8. Verification Points" if not is_ko else "8. 검증 포인트")
    if is_ko:
        rows = [
            ("Timeout", "CANoe에서 0x301 등 monitored message 송신 중단", "300 ms 이후 0x601 AliveError=1, GatewayState=DEGRADED, UART timeout log"),
            ("Unknown ID", "허용되지 않은 CAN ID 주입", "BlockedMsgCnt 증가, blocked log"),
            ("Checksum", "Byte0~Byte6 또는 Byte7 변조", "ChecksumError=1, BlockedMsgCnt 증가, checksum error log"),
            ("Alive", "특정 ECU alive counter 고정", "3회 반복 이후 AliveError=1, alive frozen log"),
            ("0x601 Period", "CANoe trace에서 0x601 관측", "100 ms 주기 송신, alive counter 0~15 wrap"),
        ]
        headers = ["Test", "Fault Injection", "Expected Result"]
    else:
        rows = [
            ("Timeout", "Stop a monitored message such as 0x301 in CANoe", "After 300 ms: 0x601 AliveError=1, GatewayState=DEGRADED, UART timeout log"),
            ("Unknown ID", "Inject an unknown CAN ID", "BlockedMsgCnt increments and blocked log is emitted"),
            ("Checksum", "Corrupt Byte0~Byte6 or Byte7", "ChecksumError=1, BlockedMsgCnt increments, checksum error log"),
            ("Alive", "Freeze a selected ECU alive counter", "After 3 repeats: AliveError=1 and alive frozen log"),
            ("0x601 Period", "Observe 0x601 in CANoe trace", "100 ms periodic transmission and alive counter wraps 0~15"),
        ]
        headers = ["Test", "Fault Injection", "Expected Result"]
    add_table(doc, headers, rows, widths=[1.25, 2.55, 3.1])

    doc.save(out)
    return out


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    seq_ko = make_sequence_diagram("ko")
    arch_ko = make_architecture_diagram("ko")
    seq_en = make_sequence_diagram("en")
    arch_en = make_architecture_diagram("en")
    ko_doc = create_doc("ko", seq_ko, arch_ko)
    en_doc = create_doc("en", seq_en, arch_en)
    print(ko_doc)
    print(en_doc)


if __name__ == "__main__":
    main()
